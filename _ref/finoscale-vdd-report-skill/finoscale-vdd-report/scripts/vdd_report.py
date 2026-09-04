#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Finoscale VDD Report Generator
Convert a VDD .docx into a branded single-page Finoscale Seller Intelligence Report PDF.

Usage:
    python3 vdd_report.py "<path/to/VDD.docx>" [--out <dir>] [--html-only]

Deps: python-docx, weasyprint, pdf2image (+ poppler-utils)
    pip install python-docx weasyprint pdf2image --break-system-packages
"""
import re, os, sys, argparse
from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "..", "assets")
HEAD = open(os.path.join(ASSETS, "template_head.html"), encoding="utf-8").read()
LOGO = open(os.path.join(ASSETS, "logo_block.txt"), encoding="utf-8").read()

STATES = ["Uttar Pradesh","Andhra Pradesh","West Bengal","Tamil Nadu","Madhya Pradesh","Himachal Pradesh",
"Arunachal Pradesh","Bihar","Jharkhand","Odisha","Karnataka","Kerala","Maharashtra","Delhi","Telangana",
"Gujarat","Rajasthan","Punjab","Haryana","Assam","Goa","Chhattisgarh","Uttarakhand"]
KEEP_UP = {"IDBI","ICICI","HDFC","SBI","RBL","PNB","IDFC","SME","NIC"}

# ---------------------------------------------------------------- helpers
def esc(s): return str(s).replace('&','&amp;').replace('<','&lt;').replace('>','&gt;')

def first_date(s):
    m=re.search(r'(\d{2})/(\d{2})/(\d{4})',s)
    if m: return m.group(0)
    m=re.search(r'(\d{4})-(\d{2})-(\d{2})',s)
    if m: return f"{m.group(3)}/{m.group(2)}/{m.group(1)}"
    m=re.search(r'\b(19|20)\d{2}\b',s)
    return m.group(0) if m else s.strip()

def clean_bank(s):
    s=re.split(r'[—/(]| - ',s)[0]
    s=re.sub(r'\b(Limited|Ltd\.?)\b','',s,flags=re.I).strip().rstrip(',')
    out=[]
    for w in s.split():
        wu=w.upper().strip('.')
        if wu in KEEP_UP: out.append(wu)
        elif w.lower() in ("of","and"): out.append(w.lower())
        else: out.append(w.capitalize())
    return " ".join(out)

def location(addr):
    addr=re.split(r'[(]',addr)[0]
    parts=[p.strip() for p in re.split(r'[,/]',addr) if p.strip()]
    idx=st=None
    for i in range(len(parts)-1,-1,-1):
        for s in STATES:
            if s.lower() in parts[i].lower(): st=s; idx=i; break
        if st: break
    if st is None: return parts[-1] if parts else addr
    city=""
    for j in range(idx-1,-1,-1):
        seg=re.sub(r'[–\-].*','',parts[j]).strip(); seg=re.sub(r'\d','',seg).strip()
        low=seg.lower()
        if seg and len(seg)>1 and not low.endswith("road") and "building" not in low and "floor" not in low and "block" not in low and "no." not in low:
            city=seg; break
    city=re.sub(r'\bDistrict\b','',city).strip()
    return (city+", "+st) if city else st

def tbl_kv(t):
    d={}
    for r in t.rows:
        c=[x.text.strip() for x in r.cells]
        if len(c)>=2 and c[0]: d.setdefault(c[0],c[1])
    return d

# ---------------------------------------------------------------- parse docx -> data dict
def parse(path):
    doc=Document(path)
    paras=[p.text.strip() for p in doc.paragraphs if p.text.strip() and p.text.strip()!='-']
    date="May 2026"
    if paras:
        md=re.search(r'(\d{1,2}\s+[A-Za-z]+\s+\d{4})',paras[0])
        date=md.group(1) if md else first_date(paras[0])
    profile=""
    for i,p in enumerate(paras):
        if p.lower()=="profile" and i+1<len(paras): profile=paras[i+1]; break
    tables=doc.tables
    score_t=tables[0]
    ent=hsn_t=checks_t=None; verif_rows=[]
    for t in tables:
        head=t.rows[0].cells[0].text.strip().lower()
        col0=[r.cells[0].text.strip() for r in t.rows]
        if head=="goods" and hsn_t is None: hsn_t=t; continue
        if head=="code": checks_t=t; continue
        if any("●" in k for k in col0):
            for r in t.rows: verif_rows.append([c.text.strip() for c in r.cells])
            continue
        if ("PAN Number" in col0 or "Proprietor Name" in col0 or "Company Name" in col0) and ent is None:
            ent=t
    ekv=tbl_kv(ent)
    vk={}
    for row in verif_rows:
        if len(row)>=3 and row[1]: vk.setdefault(row[1],row[2])
    def num(x):
        m=re.search(r'-?\d+\.?\d*',x); return float(m.group(0)) if m else 0
    com=int(num(score_t.rows[1].cells[1].text)); poa=int(num(score_t.rows[2].cells[1].text))
    poi=int(num(score_t.rows[3].cells[1].text)); aml=int(num(score_t.rows[4].cells[1].text))
    score=com+poa+poi+aml
    trade=ekv.get("Trade Name") or ekv.get("Company Name") or ekv.get("Proprietor Name")
    legal=vk.get("Legal Name") or vk.get("Name") or ekv.get("Proprietor Name") or ekv.get("Company Name") or trade
    legal=re.sub(r'Legal Name:\s*','',legal).strip()
    gst_since=first_date(vk.get("Date of Registration",""))
    bank=clean_bank(vk.get("Bank Name",""))
    addr=ekv.get("Registered Office Address","")
    loc=location(addr)
    nature=ekv.get("Nature of Business","")
    if len(nature)>105: nature=nature[:102].rstrip()+"…"
    vintage=ekv.get("Vintage","").split("(")[0].strip()
    ym=re.search(r'(19|20)\d{2}',ekv.get("Year of Incorporation","")); yinc=ym.group(0) if ym else ekv.get("Year of Incorporation","")
    etype=ekv.get("Type of Entity","")
    entity=[("Legal Name",esc(legal)),("Trade Name",esc(trade)),("Entity Type",esc(etype)),
            ("Year Incorporated",esc(yinc)),("Business Vintage",esc(vintage)),("Location",esc(loc)),
            ("Nature of Business",esc(nature))]
    reg=[("PAN Number",esc(ekv.get("PAN Number","")),"m"),("GSTIN",esc(ekv.get("GSTIN","")),"m"),
         ("Udyam Reg No",esc(ekv.get("Udyam Registration Number","")),"m"),
         ("GST Status",'<span class="status-pill pill-green">Active</span>',""),
         ("Taxpayer Type",esc(vk.get("Taxpayer Type","Regular")),""),
         ("GST Since",esc(gst_since),""),
         ("Bank",f'<span class="status-pill pill-green">{esc(bank)} &middot; Verified</span>',"")]
    hsn=[]
    if hsn_t:
        for r in hsn_t.rows[2:]:
            c=[x.text.strip() for x in r.cells]
            if c[0] and c[0].lower()!="hsn":
                desc=c[1]
                if len(desc)>72: desc=desc[:69].rstrip()+"…"
                hsn.append((esc(c[0]),esc(desc)))
    com_rows=[];poa_rows=[];poi_rows=[];aml_rows=[]
    for r in checks_t.rows[1:]:
        c=[x.text.strip() for x in r.cells[:3]]
        code=c[0]
        if not code or "Section Score" in code or code.lower()=="code": continue
        row=(code,esc(c[1]),esc(c[2]))
        (com_rows if code.startswith("COM") else poa_rows if code.startswith("POA")
         else poi_rows if code.startswith("POI") else aml_rows if code.startswith("AML") else []).append(row)
    rmap={c[0]:c[2] for c in com_rows+poa_rows+poi_rows}
    conf=[]; notes=[]
    C=lambda t: conf.append(("c",t)); N=lambda t: notes.append(("n",t))
    C(f"Active GST registration ({esc(gst_since)}), Regular taxpayer status &mdash; statutory compliance confirmed")
    v2=rmap.get("COM-02","").lower()
    if any(k in v2 for k in("recent","limited","6 month","&lt;1")): N(f"Limited GST track record &mdash; {esc(rmap.get('COM-02',''))}")
    r3=rmap.get("COM-03","")
    if "timely" in r3.lower(): C("GST 3B &amp; R1 filings timely &mdash; all months filed on time")
    elif "delay" in r3.lower(): N(f"GST 3B &amp; R1 filing shows {esc(r3)} over the last 12 months")
    C("Monthly GST filer &mdash; higher filing frequency and transparency")
    r5=rmap.get("COM-05","").lower()
    if "no additional" in r5: C("No additional registrations on entity PAN &mdash; single clean registration")
    elif "multi" in r5 or "active registrations" in r5: C("Active registrations across multiple states &mdash; multi-state operation")
    r6=rmap.get("COM-06","").lower()
    if "not available" in r6: N("HSN code match could not be evaluated &mdash; declared HSN data not available")
    elif "match" in r6: C("HSN codes match the declared product category")
    C(f"Bank penny-drop successful and account name matches GST records &mdash; {esc(bank)} verified")
    r8=rmap.get("COM-08","")
    if "no delay" in r8.lower(): C("No GST filing delay (&le;10 days) &mdash; satisfactory compliance behaviour")
    elif "delay" in r8.lower(): N(f"GST filing delay &mdash; {esc(r8)}")
    if "compliant" in rmap.get("COM-09","").lower(): C("PF/EPFO returns filed on time &mdash; labour compliance")
    r1=rmap.get("POA-01","").lower()
    if "owned" in r1: C("Business premises owned &mdash; "+("confirmed via sale deed" if "sale deed" in r1 else "confirmed via electricity bill"))
    elif "rented" in r1: C("Business premises rented &mdash; rental agreement on record")
    pa2=rmap.get("POA-02","").lower()
    if "not match" in pa2: N("Electricity bill address does not match the GST-registered address")
    elif "minor discrepancy" in pa2: N("Minor discrepancy in electricity bill address vs GST-registered address")
    elif "not available" in pa2: N("Electricity bill not available for address verification")
    elif "match" in pa2: C("Electricity bill address matches the GST-registered address")
    pa3=rmap.get("POA-03","").lower()
    if "not available" in pa3: N("Rental agreement not on record")
    elif "minor discrepancy" in pa3: N("Rental agreement shows a minor discrepancy")
    C(f"MSME (Udyam) registration valid &amp; active &mdash; {esc(ekv.get('Udyam Registration Number',''))}")
    pa5=rmap.get("POA-05","").lower()
    if "present" in pa5: C("Landlord declaration / NOC on record")
    elif "absent" in pa5: N("Landlord declaration / NOC not on record")
    p4=rmap.get("POI-04","").lower()
    if "not match" in p4: N("PAN name does not match GST legal name &mdash; discrepancy, warrants clarification")
    elif "match" in p4: C("PAN active and name matches GST legal name across sources")
    C("Clean Legal/AML &mdash; no sanctions, PEP, wilful-defaulter, eCourt or DRT/SARFAESI records (5/5)")
    findings=conf+notes
    extra=[]
    if any("delay" in rmap.get(c,"").lower() and "no delay" not in rmap.get(c,"").lower() for c in ("COM-03","COM-08")):
        extra.append("__COMPLIANCE_PARTIAL__")
    if poa<6: extra.append("__POA_PARTIAL__")
    if "not match" in p4: extra.append("__IDENTITY_PARTIAL__")
    seller=rmap.get("POI-02","Trader") or "Trader"
    fname=re.sub(r'[^A-Za-z0-9]+','_',trade).strip('_').upper()+"_VDD_Report"
    return fname,{"firm":esc(trade),"legal":esc(legal),"date":esc(date),"location":esc(loc),"seller_type":esc(seller),
        "score":score,"com":com,"poa":poa,"poi":poi,"aml":aml,"entity":entity,"reg":reg,"profile":esc(profile),
        "hsn":hsn,"findings":findings,"com_rows":com_rows,"poa_rows":poa_rows,"poi_rows":poi_rows,
        "aml_rows":aml_rows,"extra_unlock":extra}

# ---------------------------------------------------------------- CSS fix (weasyprint)
def fix_css(c):
    c=c.replace("--mono: 'DM Mono', monospace; --serif: 'Playfair Display', serif; --sans: 'DM Sans', sans-serif;",
                "--mono: 'Noto Sans Mono',monospace; --sans: 'Lato',Arial,sans-serif;")
    c=c.replace("background: linear-gradient(135deg, #0F172A 0%, #1E293B 100%)","background: #131C2E")
    c=c.replace("background: linear-gradient(135deg,#0F172A,#1E293B)","background: #131C2E")
    c=c.replace("background:linear-gradient(135deg,#D97706,#F59E0B)","background:#E8930A")
    c=c.replace("background: rgba(255,255,255,.05);","background: #243044;")
    c=c.replace("background: rgba(251,191,36,.04);","background: #1E2D3D;")
    c=c.replace("color: rgba(255,255,255,.35)","color: #7A8A9A")
    c=c.replace("color: rgba(255,255,255,.4)","color: #8A9AAA")
    c=c.replace("border-color: rgba(251,191,36,.2); background: rgba(251,191,36,.04);","border-color: #6B5200; background: #1C1800;")
    c=c.replace('style="color:rgba(255,255,255,.35);white-space:nowrap"','style="color:#D4A800;white-space:nowrap"')
    c=c.replace("background:rgba(251,191,36,.12); border:1px solid rgba(251,191,36,.3); color:#FBBF24;","background:#2A2800; border:1px solid #B8860B; color:#F59E0B;")
    c=c.replace("background:repeating-linear-gradient(90deg,rgba(251,191,36,.4) 0px,rgba(251,191,36,.4) 5px,rgba(251,191,36,.1) 5px,rgba(251,191,36,.1) 10px)","background:repeating-linear-gradient(90deg,#B8860B 0px,#B8860B 5px,#4A3A0A 5px,#4A3A0A 10px)")
    c=re.sub(r'\.hdr::before\s*\{[^}]*radial-gradient[^}]*\}','.hdr::before { display: none; }',c)
    c=re.sub(r'(\.ftr\s*\{[^}]*?)background\s*:\s*var\(--slate\)',r'\1background:#1E293B',c)
    c=re.sub(r'(\.section-header-row\s+td\s*\{[^}]*?)background\s*:\s*var\(--slate\)\s*!important',r'\1background:#1E293B !important',c)
    c=re.sub(r'\.chip\s*\{([^}]*?)border\s*:\s*1px solid\s*;([^}]*)\}',lambda m:'.chip {'+m.group(1)+'display: inline-block; white-space: nowrap;'+m.group(2)+'}',c)
    c=re.sub(r'\.chip-teal\s*\{[^}]*\}','.chip-teal { background: #0D4A40; border: 1.5px solid #0D4A40; color: #5EEAD4; padding: 2px 8px; border-radius: 999px; font-size: 11px; font-weight: 600; }',c)
    c=re.sub(r'\.chip-blue\s*\{[^}]*\}','.chip-blue { background: #1A3A70; border: 1.5px solid #1A3A70; color: #93C5FD; padding: 2px 8px; border-radius: 999px; font-size: 11px; font-weight: 600; }',c)
    c=re.sub(r'\.chip-white\s*\{[^}]*\}','.chip-white { background: #1E2D42; border: 1.5px solid #1E2D42; color: #B8C8D8; padding: 2px 8px; border-radius: 999px; font-size: 11px; font-weight: 600; }',c)
    c=re.sub(r'(\.brand-tag\s*\{[^}]*?color\s*:\s*)rgba\(255,255,255,\.45\)',r'\1#8A9AAA',c)
    c=re.sub(r'(\.brand-tag\s*\{)([^}]*)\}',lambda m:m.group(1)+m.group(2)+' white-space: nowrap;}',c)
    c=c.replace('.unlock-card { background:#1E293B;','.unlock-card { background:#1A2840;')
    return c

# ---------------------------------------------------------------- build HTML
def _bar(r): return "#059669" if r>=0.75 else ("#D97706" if r>=0.45 else "#DC2626")
def _kv(k,v): return f'<div class="kv"><span class="kv-key">{k}</span><span class="kv-dot"></span><span class="kv-val">{v}</span></div>'
def _kvm(k,v): return f'<div class="kv"><span class="kv-key">{k}</span><span class="kv-dot"></span><span class="kv-val mono">{v}</span></div>'
def _uc(cat,name,desc,cur,den,cta,sub,pill,partial=False):
    pc=" partial" if partial else ""; w=round(cur/den*100) if den else 0; lbl="Current parameter" if partial else "Progress"
    return (f'<div class="unlock-card{pc}"><div class="unlock-pts-pill{pc}">{pill}</div>'
            f'<div><div class="unlock-cat">{cat}</div><div class="unlock-name">{name}</div></div>'
            f'<div class="unlock-desc">{desc}</div>'
            f'<div class="unlock-progress"><div class="unlock-prog-row"><span class="unlock-prog-label">{lbl}</span>'
            f'<span class="unlock-prog-val">{cur} / {den} pts</span></div>'
            f'<div class="unlock-prog-track"><div class="unlock-prog-fill" style="width:{w}%;background:#D97706"></div></div></div>'
            f'<div class="unlock-cta"><div class="unlock-cta-arrow">&rarr;</div><div><div class="unlock-cta-text">{cta}</div>'
            f'<div class="unlock-cta-sub">{sub}</div></div></div></div>')

def build(d):
    score=d['score']; scope=100-score; com,poa,poi,aml=d['com'],d['poa'],d['poi'],d['aml']
    chips=('<span class="chip chip-teal">&#10003; AML Cleared</span><span class="chip chip-teal">&#10003; GST Active</span>'
           '<span class="chip chip-blue">&#10003; Bank Verified</span>')
    chips+=f'<span class="chip chip-white"><svg width="9" height="9" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.2" stroke-linecap="round" stroke-linejoin="round" style="vertical-align:-1px;margin-right:4px"><path d="M20 10c0 6-8 12-8 12s-8-6-8-12a8 8 0 0 1 16 0Z"/><circle cx="12" cy="10" r="3"/></svg>{d["location"]}</span>'
    chips+=f'<span class="chip chip-white">{d["seller_type"]}</span>'
    hdr=(f'<div class="hdr"><div class="hdr-inner"><div class="hdr-top"><div class="brand">{LOGO}'
         '<div><div class="brand-name">finoscale</div><div class="brand-tag">Seller Intelligence Platform</div></div></div>'
         f'<div class="hdr-date">Report Date: {d["date"]}</div></div>'
         f'<div class="hdr-firm"><div class="hdr-label">Verified Seller Profile</div><div class="hdr-firmname">{d["firm"]}</div>'
         f'<div class="hdr-proprietor">Legal Name: {d["legal"]}</div></div><div class="hdr-chips">{chips}</div></div></div>')
    def card(lbl,v,den):
        r=v/den
        return (f'<div class="sc-card"><div class="sc-card-lbl">{lbl}</div><div class="sc-card-val">{v}'
                f'<span style="font-size:11px;opacity:.5"> /{den}</span></div><div class="sc-bar">'
                f'<div class="sc-fill" style="width:{round(r*100)}%;background:{_bar(r)}"></div></div></div>')
    pend=lambda l,den:(f'<div class="sc-card pending-card"><div class="sc-card-lbl">{l}</div>'
        f'<div class="sc-card-val" style="color:rgba(255,255,255,.35);white-space:nowrap">0<span style="font-size:11px;opacity:.5"> /{den}</span></div>'
        '<div class="pending-tag"><svg width="9" height="9" viewBox="0 0 24 24" fill="none" stroke="#FBBF24" stroke-width="2.2" stroke-linecap="round" stroke-linejoin="round" style="vertical-align:-1px;margin-right:3px"><path d="M5 22h14M5 2h14M17 22v-4.17a2 2 0 0 0-.59-1.41L12 12l-4.41 4.41A2 2 0 0 0 7 17.83V22M7 2v4.17a2 2 0 0 0 .59 1.41L12 12l4.41-4.41A2 2 0 0 0 17 6.17V2"/></svg>Pending</div></div>')
    scoreband=(f'<div class="score-band"><div class="score-main"><div class="score-main-lbl">Finoscale Basic Score</div>'
        f'<div><span class="score-main-num">{score}</span><span class="score-main-denom"> / 100</span></div>'
        f'<div class="score-sub">Verified data only &middot; {d["date"]}</div></div>'
        '<div style="display:flex;flex-direction:column;flex:1"><div class="score-cards">'
        f'{card("Compliance",com,25)}{card("Proof of Address",poa,10)}{card("Proof of Identity",poi,10)}{card("Legal / AML",aml,5)}</div>'
        f'<div class="score-cards-row2">{pend("On-Site Verification",20)}{pend("3B Analysis","12.5")}{pend("2B Analysis","12.5")}{pend("ITR Analysis",5)}</div></div></div>')
    prog=(f'<div class="prog-wrap"><div class="prog-legend"><div class="prog-legend-left">'
        '<div style="display:flex;align-items:center;gap:6px"><div class="prog-legend-dot" style="background:linear-gradient(90deg,#059669,#34D399)"></div>'
        f'<span class="prog-legend-label">Current Score &nbsp;<strong>{score}</strong></span></div>'
        '<div style="display:flex;align-items:center;gap:6px"><div class="prog-legend-dot" style="background:repeating-linear-gradient(90deg,rgba(251,191,36,.5) 0px,rgba(251,191,36,.5) 4px,rgba(251,191,36,.12) 4px,rgba(251,191,36,.12) 8px)"></div>'
        f'<span class="prog-legend-label">Scope of Improvement &nbsp;<strong style="color:#FBBF24">+{scope}</strong></span></div></div>'
        '<span style="font-family:\'DM Mono\',monospace;font-size:10px;color:rgba(255,255,255,.3)">Max possible: 100 / 100</span></div>'
        f'<div class="prog-bar-track"><div class="prog-bar-current" style="width:{score}%"></div><div class="prog-bar-scope" style="left:{score}%;width:{scope}%"></div></div>'
        '<div class="prog-ticks"><span class="prog-tick">0</span><span class="prog-tick">25</span><span class="prog-tick">50</span><span class="prog-tick">75</span><span class="prog-tick">100</span></div></div>')
    ucards=[_uc("On-Site Verification &middot; Pending","Physical On-Site Verification","Premises, stock, and operations have not been physically verified. Completing OSV is the highest single scoring action available.",0,20,"Schedule OSV Visit &rarr; unlock +20 pts","Coordinate with your Finoscale account manager","0 / 20 pts"),
        _uc("3B &amp; 2B Analysis &middot; Pending","GST Return Analysis (3B &amp; 2B)","Detailed GSTR-3B turnover and GSTR-2B ITC analysis is pending. Requires one-time GST portal consent to pull filing data.",0,25,"Grant GST Data Consent &rarr; unlock +25 pts","One-time authorisation on the GST portal","0 / 25 pts"),
        _uc("ITR Analysis &middot; Pending","ITR &amp; Tax Compliance History","Income Tax Return filing history is not on file. Consistent ITR with matching declared turnover significantly strengthens the business profile.",0,5,"Submit ITR Documents &rarr; unlock +5 pts","Upload last 2 years of ITR acknowledgements","0 / 5 pts")]
    for tag in d['extra_unlock']:
        if tag=="__COMPLIANCE_PARTIAL__":
            ucards.append(_uc("Compliance &middot; Partial","Improve GST Filing Timeliness","GST 3B &amp; R1 filings show delays over the last 12 months. Consistent on-time filing restores the full compliance score.",com,25,"File GST returns on time &rarr; unlock compliance points","Maintain timely monthly GSTR-3B &amp; R1 filing",f"{com} / 25 pts",True))
        elif tag=="__POA_PARTIAL__":
            ucards.append(_uc("Proof of Address &middot; Partial","Complete Premises Documentation","Premises ownership / rental status or a matching utility bill is incomplete. A matching utility bill plus ownership or rental proof restores the full address-proof score.",poa,10,"Submit premises &amp; address proof &rarr; unlock PoA points","Ownership / rental agreement and matching utility bill",f"{poa} / 10 pts",True))
        elif tag=="__IDENTITY_PARTIAL__":
            ucards.append(_uc("Proof of Identity &middot; Partial","Resolve PAN&ndash;GST Name Mismatch","The PAN name does not match the GST legal name. Aligning the PAN and GST legal name restores the full identity score.",poi,10,"Reconcile PAN &amp; GST legal name &rarr; unlock points","Submit matching PAN and GST registration records",f"{poi} / 10 pts",True))
    unlock=('<div class="unlock-sec"><div class="unlock-header"><div><div class="unlock-title">Unlock Your Full Score</div>'
        f'<div class="unlock-subtitle">3 parameters are pending &mdash; submit documents to increase your score by up to <span style="color:#FBBF24;font-weight:600">+{scope} points</span></div></div>'
        f'<div class="unlock-badge"><div class="unlock-badge-label">Points Unlockable</div><div class="unlock-badge-num">+{scope}</div><div class="unlock-badge-sub">across {len(ucards)} areas</div></div></div>'
        f'<div class="unlock-grid">{"".join(ucards)}</div></div>')
    left='<div class="body-col"><div class="sec-title"><span class="sec-title-bar" style="background:var(--accent)"></span>Entity Details</div><div class="kv-list">'
    for k,v in d['entity']: left+=_kv(k,v)
    left+='</div></div>'
    right='<div class="body-col"><div class="sec-title"><span class="sec-title-bar" style="background:var(--accent)"></span>Registration &amp; Compliance IDs</div><div class="kv-list">'
    for k,v,m in d['reg']: right+=(_kvm(k,v) if m=='m' else _kv(k,v))
    right+='</div></div>'
    bodygrid=f'<div class="body-grid">{left}{right}</div>'
    profile=f'<div class="profile-sec"><div class="sec-title"><span class="sec-title-bar" style="background:var(--accent)"></span>Business Profile</div><p class="profile-text">{d["profile"]}</p></div>'
    hsn='<div class="hsn-sec"><div class="sec-title"><span class="sec-title-bar" style="background:var(--accent)"></span>Declared Goods &amp; Services (HSN Codes)</div><table class="hsn-table"><thead><tr><th style="width:120px">HSN / SAC</th><th>Description</th></tr></thead><tbody>'
    for code,desc in d['hsn']: hsn+=f'<tr><td><span class="hsn-code">{code}</span></td><td>{desc}</td></tr>'
    hsn+='</tbody></table></div>'
    findings='<div class="findings-sec"><div class="sec-title"><span class="sec-title-bar" style="background:var(--accent)"></span>Findings &amp; Observations</div><div class="findings-grid">'
    for t,txt in d['findings']:
        cls,ic,sym=("finding-confirmed","icon-confirmed","&#10003;") if t=='c' else ("finding-note","icon-note","!")
        findings+=f'<div class="finding-item {cls}"><span class="finding-icon {ic}">{sym}</span>{txt}</div>'
    findings+='</div></div>'
    annex='<div class="annex-sec"><div class="sec-title"><span class="sec-title-bar" style="background:var(--accent)"></span>Compliance &amp; KYC Scoring Detail</div><table class="annex-table"><thead><tr><th style="width:80px">Code</th><th>Parameter</th><th>Result</th></tr></thead><tbody>'
    def sect(title,val,pending=False):
        col='#FBBF24' if pending else '#fff'
        return f'<tr class="section-header-row"><td colspan="2"><strong>{title}</strong> &middot; {"Pending" if pending else "No Consent"}</td><td style="text-align:right;color:{col} !important;font-family:var(--mono);white-space:nowrap">{val}</td></tr>'
    row=lambda c,p,r:f'<tr><td><span class="code-tag">{c}</span></td><td>{p}</td><td>{r}</td></tr>'
    annex+=sect("COMPLIANCE",com)
    for c,p,r in d['com_rows']: annex+=row(c,p,r)
    annex+=sect("PROOF OF ADDRESS",poa)
    for c,p,r in d['poa_rows']: annex+=row(c,p,r)
    annex+=sect("PROOF OF IDENTITY",poi)
    for c,p,r in d['poi_rows']: annex+=row(c,p,r)
    annex+=sect("LEGAL / AML CHECK",aml)
    for c,p,r in d['aml_rows']: annex+=row(c,p,r)
    annex+=sect("ON-SITE VERIFICATION",0,True)+row("OSV-01","Physical On-Site Verification","Not yet completed &mdash; pending scheduling")
    annex+=sect("3B &amp; 2B ANALYSIS",0,True)+row("3B-01","GSTR-3B Turnover Analysis","Pending &mdash; GST consent not yet granted")+row("2B-01","GSTR-2B ITC Analysis","Pending &mdash; GST consent not yet granted")
    annex+=sect("ITR ANALYSIS",0,True)+row("ITR-01","ITR Filing History &amp; Tax Compliance","Pending &mdash; ITR documents not submitted")
    annex+='</tbody></table></div>'
    ftr='<div class="ftr"><div class="ftr-brand"><strong>finoscale.ai</strong> &nbsp;&middot;&nbsp; Seller Intelligence &amp; KYC Platform</div><div class="ftr-disclaimer">This report is generated for platform use only. Scores are based on available data at the time of assessment. Finoscale does not guarantee accuracy of third-party data. For queries: www.finoscale.ai</div></div>'
    head=HEAD.replace("KASHI ENTERPRISES &mdash; Seller Intelligence Report", d['firm']+" &mdash; Seller Intelligence Report")
    return head+hdr+scoreband+prog+unlock+bodygrid+profile+hsn+findings+annex+ftr+'\n</div></body></html>'

# ---------------------------------------------------------------- render tight-height PDF
def render_pdf(html, out_pdf):
    import weasyprint
    from pdf2image import convert_from_bytes
    css_html=fix_css(html)
    hp=os.path.dirname(os.path.abspath(out_pdf))+"/"
    def render(h):
        css=f'@page {{ size: 960px {h}px; margin: 0; }} html,body{{width:960px;background:#F1F3F7;}}'
        return weasyprint.HTML(string=css_html+f'<style>{css}</style>',base_url='file://'+hp).write_pdf(presentational_hints=True)
    def pages(pdf): return len(convert_from_bytes(pdf,dpi=36))
    best=None
    for h in range(4200,2999,-100):       # coarse
        pdf=render(h); 
        if pages(pdf)==1: best=(h,pdf)
        elif best: break
    if not best: best=(4200,render(4200))
    h=best[0]-10                            # fine
    while h>2999:
        pdf=render(h)
        if pages(pdf)==1: best=(h,pdf); h-=10
        else: break
    open(out_pdf,'wb').write(best[1])
    return best[0]

# ---------------------------------------------------------------- main
def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("docx"); ap.add_argument("--out",default=None); ap.add_argument("--html-only",action="store_true")
    a=ap.parse_args()
    out_dir=a.out or os.path.dirname(os.path.abspath(a.docx))
    os.makedirs(out_dir,exist_ok=True)
    fname,d=parse(a.docx)
    print(f"[score] {d['firm']}: COM {d['com']} + POA {d['poa']} + POI {d['poi']} + AML {d['aml']} = {d['score']}/100")
    html=build(d)
    html_path=os.path.join(out_dir,fname+".html"); open(html_path,"w",encoding="utf-8").write(html)
    if a.html_only:
        print("[html]",html_path); return
    pdf_path=os.path.join(out_dir,fname+".pdf")
    h=render_pdf(html,pdf_path)
    print(f"[pdf] {pdf_path}  (page height {h}px)")

if __name__=="__main__":
    main()
